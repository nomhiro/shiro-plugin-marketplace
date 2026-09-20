# -*- coding: utf-8 -*-
"""make-lecture-slides 用ヘルパー。template.pptx を継承して動画用スライドを組む Deck クラス。

使い方:
    import sys; sys.path.insert(0, r"<skill>/scripts")
    from pptx_helpers import Deck, NAVY, BLUE, SLATE, LTBLUE, ORANGE, LTOR, GREEN, LTGRN, GRAY, LTGRAY, WHITE, INK
    d = Deck(r"<repo>/template.pptx")
    d.title_slide("タイトル", [{"text":"サブ","size":18,"bold":True,"color":SLATE}])
    d.text_slide("ゴール", [("項目1",0),("補足",1)])
    s = d.diagram("章タイトル"); d.box(s,0.92,1.6,5,2,LTBLUE,line=BLUE); d.textbox(s,1.1,1.8,4.6,1.6,[{"text":"…","size":14}])
    d.table(d.diagram("表"), rows, 0.92,1.7,11.49,[3,4,4.49])
    d.save(r"<repo>/lectures/01_plan_manage/L1-1-1_座学_foundry-model-map.pptx")
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

EMU_IN = 914400
FONT = "BIZ UDPGothic"

# テンプレ theme accent ベースのパレット
NAVY = RGBColor(0x1F, 0x38, 0x64)
SLATE = RGBColor(0x44, 0x54, 0x6A)
BLUE = RGBColor(0x44, 0x72, 0xC4)
LTBLUE = RGBColor(0xD9, 0xE2, 0xF3)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
LTOR = RGBColor(0xFB, 0xE5, 0xD6)
GREEN = RGBColor(0x54, 0x82, 0x35)
LTGRN = RGBColor(0xE2, 0xEF, 0xDA)
GRAY = RGBColor(0x7F, 0x7F, 0x7F)
LTGRAY = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x26, 0x26, 0x26)
CODE_BG = RGBColor(0x1E, 0x1E, 0x1E)
CODE_FG = RGBColor(0xE6, 0xE6, 0xE6)

_RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def IN(v):
    return Emu(int(v * EMU_IN))


class Deck:
    """template.pptx を開き、既定スライドを除去して構築するラッパ。
    レイアウト index: 0=セクションタイトル, 1=タイトルと説明, 2=BLANK（テンプレ依存）。"""

    def __init__(self, template_path, font=FONT, layout_section=0, layout_titlebody=1):
        self.prs = Presentation(template_path)
        self.font = font
        lst = self.prs.slides._sldIdLst
        for s in list(lst):
            r = s.get(_RID)
            if r:
                self.prs.part.drop_rel(r)
            lst.remove(s)
        self.W = self.prs.slide_width / EMU_IN
        self.H = self.prs.slide_height / EMU_IN
        self.CL, self.CR = 0.92, 12.41
        self.CW = self.CR - self.CL
        self._lay_section = layout_section
        self._lay_titlebody = layout_titlebody

    # ---- slide-level ----
    def _set_title(self, s, text):
        t = s.shapes.title
        t.text = text
        for p in t.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True
        return t

    def _drop_body(self, s):
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx == 1:
                ph._element.getparent().remove(ph._element)

    def title_slide(self, title, sublines, accent=BLUE):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[self._lay_section])
        self._set_title(s, title)
        self.box(s, self.CL, 4.6, 3.2, 0.10, accent, rounded=False)
        self.textbox(s, self.CL, 4.8, self.CW, 1.7, sublines)
        return s

    def text_slide(self, title, bullets):
        """bullets: list of (text, level)。level0=見出し色, 1=本文。"""
        s = self.prs.slides.add_slide(self.prs.slide_layouts[self._lay_titlebody])
        self._set_title(s, title)
        tf = s.placeholders[1].text_frame
        tf.word_wrap = True
        for i, (txt, lvl) in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.level = lvl
            r = p.add_run(); r.text = txt
            r.font.size = Pt(21 if lvl == 0 else 17)
            if lvl == 0:
                r.font.bold = True; r.font.color.rgb = NAVY
            else:
                r.font.color.rgb = INK
        return s

    def diagram(self, title):
        """タイトル＋区切り線（テンプレ継承）だけのスライドを返す。本体は box/textbox/table で組む。"""
        s = self.prs.slides.add_slide(self.prs.slide_layouts[self._lay_titlebody])
        self._set_title(s, title)
        self._drop_body(s)
        return s

    # ---- shape-level (slide を第1引数に) ----
    def box(self, s, x, y, w, h, fill, line=None, rounded=True, line_w=1.25):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                IN(x), IN(y), IN(w), IN(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line; sh.line.width = Pt(line_w)
        sh.shadow.inherit = False
        return sh

    def put_lines(self, sh, lines, pad=0.1, anchor=MSO_ANCHOR.MIDDLE):
        """既存シェイプ(sh)にテキスト行を入れる。lines: list of dict(text,size,bold,color,align,space_after)。"""
        tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, IN(pad))
        self._fill(tf, lines)
        return sh

    def textbox(self, s, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, IN(0.02))
        self._fill(tf, lines)
        return tb

    @staticmethod
    def _fill(tf, lines):
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ln.get("align", PP_ALIGN.LEFT)
            if "space_after" in ln:
                p.space_after = Pt(ln["space_after"])
            r = p.add_run(); r.text = ln["text"]
            r.font.size = Pt(ln.get("size", 14))
            r.font.bold = ln.get("bold", False)
            r.font.color.rgb = ln.get("color", INK)

    def arrow(self, s, x, y, w, h, kind=MSO_SHAPE.RIGHT_ARROW, fill=None):
        sh = s.shapes.add_shape(kind, IN(x), IN(y), IN(w), IN(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = fill or BLUE
        sh.line.fill.background(); sh.shadow.inherit = False
        return sh

    def numbered(self, s, x, y, n, color=BLUE, d=0.45):
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, IN(x), IN(y), IN(d), IN(d))
        c.fill.solid(); c.fill.fore_color.rgb = color; c.line.fill.background(); c.shadow.inherit = False
        self.put_lines(c, [{"text": str(n), "size": 18, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}], pad=0.0)
        return c

    def table(self, s, rows, x, y, w, colw, h=None, header_fill=NAVY, fs=14):
        """rows: list of tuples（1行目=ヘッダ）。colw: 各列幅(inch)。h省略時は行数から概算。"""
        if h is None:
            h = 0.55 * len(rows)
        t = s.shapes.add_table(len(rows), len(rows[0]), IN(x), IN(y), IN(w), IN(h)).table
        for i, cwi in enumerate(colw):
            t.columns[i].width = IN(cwi)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.margin_left = IN(0.12); cell.margin_right = IN(0.08)
                cell.margin_top = IN(0.04); cell.margin_bottom = IN(0.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = cell.text_frame.paragraphs[0]; r = p.add_run(); r.text = val
                if ri == 0:
                    r.font.bold = True; r.font.size = Pt(fs + 1); r.font.color.rgb = WHITE
                    cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
                else:
                    r.font.size = Pt(fs); r.font.color.rgb = INK; r.font.bold = (ci == 0)
                    cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if ri % 2 else LTGRAY
        return t

    def code_box(self, s, x, y, w, h, code, size=13):
        """暗背景＋等幅(Consolas)のコードブロック。code は改行入り文字列。"""
        self.box(s, x, y, w, h, CODE_BG)
        tb = s.shapes.add_textbox(IN(x + 0.2), IN(y + 0.15), IN(w - 0.4), IN(h - 0.3))
        tf = tb.text_frame; tf.word_wrap = True
        for i, line in enumerate(code.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run(); r.text = line if line else " "
            r.font.size = Pt(size); r.font.name = "Consolas"; r.font.color.rgb = CODE_FG
        return tb

    # ---- finalize ----
    def _apply_font(self):
        for sl in self.prs.slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            self._set_run_font(r)
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        for cell in row.cells:
                            for p in cell.text_frame.paragraphs:
                                for r in p.runs:
                                    self._set_run_font(r)

    def _set_run_font(self, run):
        # Consolas（コード）は等幅のまま残す
        if run.font.name == "Consolas":
            return
        run.font.name = self.font
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {}); rPr.append(el)
            el.set("typeface", self.font)

    def save(self, out):
        self._apply_font()
        self.prs.save(out)
        return out
